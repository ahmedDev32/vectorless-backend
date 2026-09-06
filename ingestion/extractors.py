"""Per-file-type text extraction functions.

Each function takes a file path and returns the raw extracted text.
Kept dependency-light: only libraries already installed are used.
"""

import csv


def extract_txt(path: str) -> str:
    with open(path, encoding='utf-8', errors='replace') as f:
        return f.read()


def extract_csv(path: str) -> str:
    with open(path, newline='', encoding='utf-8', errors='replace') as f:
        reader = csv.reader(f)
        return '\n'.join(', '.join(row) for row in reader)


def extract_pdf(path: str) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    return '\n\n'.join(page.extract_text() or '' for page in reader.pages)


def extract_docx(path: str) -> str:
    import docx

    document = docx.Document(path)
    return '\n'.join(paragraph.text for paragraph in document.paragraphs)


def extract_pptx(path: str) -> str:
    from pptx import Presentation

    presentation = Presentation(path)
    lines = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines.append(shape.text_frame.text)
    return '\n'.join(lines)
